"""
Generate a before/after progress Word document.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

doc = Document()

section = doc.sections[0]
section.top_margin    = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin   = Cm(2.5)
section.right_margin  = Cm(2.5)

# ── Helpers ───────────────────────────────────────────────────────────────────
def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False, color=None, size=10, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = align
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

def add_table(headers, rows, col_widths, header_color="1F4E79",
              before_col=1, after_col=2):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_bg(cell, header_color)
        set_cell_text(cell, h, bold=True, color=(255,255,255),
                      size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Data rows
    for ri, row_data in enumerate(rows):
        row = table.rows[ri + 1]
        row_bg = "F2F2F2" if ri % 2 == 0 else "FFFFFF"
        for ci, val in enumerate(row_data):
            cell = row.cells[ci]
            # Colour before/after columns distinctly
            if ci == before_col:
                set_cell_bg(cell, "FDECEA")   # light red
                set_cell_text(cell, str(val), size=10)
            elif ci == after_col:
                set_cell_bg(cell, "E8F5E9")   # light green
                set_cell_text(cell, str(val), size=10)
            else:
                set_cell_bg(cell, row_bg)
                set_cell_text(cell, str(val), bold=(ci == 0), size=10)

    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)

    doc.add_paragraph()
    return table

def h1(text):
    p = doc.add_heading(text, level=1)
    return p

def h2(text):
    p = doc.add_heading(text, level=2)
    return p

def body(text):
    p = doc.add_paragraph(text)
    p.style.font.size = Pt(11)
    return p

# ── Title ─────────────────────────────────────────────────────────────────────
title = doc.add_heading("3D Reconstruction Pipeline — Before & After", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

sub = doc.add_paragraph("Quality Improvements — Week of April 2026")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].font.size = Pt(13)
sub.runs[0].font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

doc.add_paragraph()

# ── Legend ────────────────────────────────────────────────────────────────────
legend_para = doc.add_paragraph()
legend_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
r1 = legend_para.add_run("  Before  ")
r1.font.size = Pt(10); r1.bold = True
r1.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

r2 = legend_para.add_run("     ")

r3 = legend_para.add_run("  After  ")
r3.font.size = Pt(10); r3.bold = True
r3.font.color.rgb = RGBColor(0x37, 0x56, 0x23)

doc.add_paragraph()

# ═══════════════════════════════════════════════════════════════════════════════
# TABLE 1 — Output Quality
# ═══════════════════════════════════════════════════════════════════════════════
h1("1. Gaussian Splat Output Quality")

add_table(
    headers=["Metric", "Before (V1 — Sparse Init)", "After (V3 — MVS 500k Init)"],
    rows=[
        ["Total Gaussians",             "4,998,420",         "4,839,487"],
        ["Ghost Gaussians (opacity<0.1)","36.2%  — 1.81M",   "34.0%  — 1.64M"],
        ["Solid Gaussians (opacity>0.9)","25.8%  — 1.29M",   "27.1%  — 1.31M  ↑"],
        ["Mean opacity (sigmoid)",       "0.405",             "0.433  ↑"],
        ["Median opacity",               "0.209",             "0.247  ↑"],
        ["Gaussian initialisation",      "37,710 sparse pts", "500,000 MVS dense pts"],
        ["File size (full splat)",       "1.20 GB",           "1.14 GB"],
        ["File size (after pruning)",    "659 MB",            "668 MB"],
        ["Gaussians after pruning",      "2,786,436",         "2,824,753  ↑"],
        ["Pruning applied",              "Manual (one-off)",  "Automatic on every run"],
    ],
    col_widths=[2.4, 2.3, 2.8],
    header_color="1F4E79",
)

# ═══════════════════════════════════════════════════════════════════════════════
# TABLE 2 — Initialisation
# ═══════════════════════════════════════════════════════════════════════════════
h1("2. Gaussian Initialisation")

add_table(
    headers=["Aspect", "Before", "After"],
    rows=[
        ["Source of init points",   "COLMAP sparse SfM only",              "COLMAP MVS dense reconstruction"],
        ["Number of init points",   "37,710",                               "500,000 (subsampled from 5M)"],
        ["Init file",               "sparse_init.ply  (< 1 MB)",           "fused_sub.ply  (13 MB)"],
        ["Coverage quality",        "Keypoint-only — misses flat surfaces", "Dense per-pixel — covers walls, floors, furniture"],
        ["GPU memory at start",     "~100 MB (trivial)",                    "~1.2 GB (manageable, stable)"],
        ["MVS runtime cost",        "Not run",                              "~47 min (one-time, result reused)"],
    ],
    col_widths=[2.2, 2.3, 3.0],
    header_color="1F4E79",
)

# ═══════════════════════════════════════════════════════════════════════════════
# TABLE 3 — Pipeline Improvements
# ═══════════════════════════════════════════════════════════════════════════════
h1("3. Pipeline Capabilities")

add_table(
    headers=["Feature", "Before", "After"],
    rows=[
        ["Output versioning",          "Always overwrites splat.ply",                  "Auto-versioned: splat_v3_mvs_500k.ply etc."],
        ["Ghost Gaussian pruning",     "Not implemented",                              "Automatic post-processing on every run"],
        ["Quality analysis",           "Visual inspection only",                       "Statistical analysis script (opacity, scale, outliers)"],
        ["HEIC support",               "JPG/PNG only",                                 "Full HEIC conversion (RealityScan native)"],
        ["Best model selection",       "Hardcoded sparse/0",                           "Auto-selects largest points3D.bin"],
        ["Idempotency",                "Re-ran everything from scratch",               "Skips completed stages (COLMAP, downscale, etc.)"],
        ["fused.ply preservation",     "Wiped by COLMAP undistorter on re-run",        "Backup/restore logic protects it"],
        ["LPIPS model cache",          "Re-downloaded 233 MB alexnet every run",       "Cached in workspace/torch_cache — downloaded once"],
        ["torch.load compatibility",   "Crashed on export and resume",                 "Patched eval_utils.py + trainer.py automatically"],
        ["Training resume",            "Not possible",                                 "--resume flag (note: 3DGS resume is fragile by design)"],
        ["PLY subsampling",            "Not implemented",                              "subsample_ply.py — reduces init size for stability"],
    ],
    col_widths=[2.2, 2.3, 3.0],
    header_color="1F4E79",
)

# ═══════════════════════════════════════════════════════════════════════════════
# TABLE 4 — Stability
# ═══════════════════════════════════════════════════════════════════════════════
h1("4. System Stability")

add_table(
    headers=["Issue", "Before", "After"],
    rows=[
        ["Windows BSOD (HYPERVISOR_ERROR)", "Crashed 3× during MVS-init training",        "HAGS disabled, TDR timeout extended — stable"],
        ["GPU TDR timeout",                 "Default 2 sec — killed driver mid-compute",   "Extended to 60 sec via registry"],
        ["WSL2 memory management",          "Default settings",                             "autoMemoryReclaim=gradual in .wslconfig"],
        ["MVS full init (5M pts)",          "Crashed GPU/Hyper-V at step ~4,000",          "Solved by subsampling to 500k before training"],
        ["3DGS checkpoint resume",          "Not implemented",                              "Added --resume flag (fresh start preferred due to ADC)"],
        ["Docker model cache",              "233 MB download on every training run",        "-e TORCH_HOME=/workspace/torch_cache on all runs"],
    ],
    col_widths=[2.5, 2.3, 2.7],
    header_color="7F3F00",
)

# ═══════════════════════════════════════════════════════════════════════════════
# TABLE 5 — What's Next
# ═══════════════════════════════════════════════════════════════════════════════
h1("5. Next Steps")

add_table(
    headers=["Task", "Expected Outcome", "Status"],
    rows=[
        ["Mesh export (Poisson)",           "Physics-ready mesh for VR object interaction",    "Ready to run"],
        ["2DGS method",                     "Surface-aligned Gaussians → better mesh quality", "Needs Docker image update"],
        ["Ubuntu dual boot",                "Eliminate BSOD risk, native CUDA performance",    "Planned"],
        ["Unity VR integration",            "Gaussian splat visible in Quest 3 headset",       "Pending mesh export"],
        ["Per-object segmentation (SAM3D)", "Individual physics bodies per room object",       "Future work"],
    ],
    col_widths=[2.5, 3.0, 2.0],
    header_color="1F4E79",
    before_col=99, after_col=99,   # no colour on these columns
)

# ── Save ──────────────────────────────────────────────────────────────────────
out = r"C:\Users\mgallai\Projects\3d_automated\reconstruction_project\Pipeline_Before_After.docx"
doc.save(out)
print(f"Saved: {out}")
